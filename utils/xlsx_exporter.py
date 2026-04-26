"""
분석 결과를 Excel(xlsx) 형식으로 내보내는 모듈.
openpyxl을 사용하여 상세 취약점 내역을 엑셀 파일로 저장합니다.
"""
import os
from datetime import datetime
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def export_xlsx(findings: List[Dict[str, Any]], output_path: str) -> str:
    """
    취약점 리스트를 Excel 파일로 저장합니다.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Vulnerability Findings"

    # 헤더 정의
    headers = [
        "심각도", "취약점 제목", "발생 위치(파일)", "취약점 유형", 
        "ISMS-P 위반사항", "상세 설명", "수정 가이드", "판정(TP/FP)"
    ]
    
    # 헤더 스타일 설정
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'), 
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )

    for col_num, column_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = border

    # 데이터 작성
    for row_num, finding in enumerate(findings, 2):
        is_tp = finding.get("is_true_positive", True)
        status = "True Positive" if is_tp else "False Positive"
        
        # 파일 경로 추출
        file_path = finding.get("file_path", finding.get("original_file", "N/A"))
        
        row_data = [
            finding.get("severity", "INFO"),
            finding.get("title", "제목 없음"),
            file_path,
            finding.get("vulnerability_type", finding.get("rule_id", "N/A")),
            finding.get("isms_p_violation", "해당 없음"),
            finding.get("description", "설명 없음"),
            finding.get("remediation_description", "가이드 없음"),
            status
        ]
        
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            
            # 심각도별 색상 강조 (선택 사항)
            if col_num == 1:
                sev = str(value).upper()
                if sev == "CRITICAL":
                    cell.font = Font(color="FF0000", bold=True)
                elif sev == "HIGH":
                    cell.font = Font(color="FF8C00", bold=True)
                elif sev == "MEDIUM":
                    cell.font = Font(color="0000FF")

    # 컬럼 너비 조정
    column_widths = [12, 40, 40, 20, 40, 60, 60, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    wb.save(output_path)
    return output_path
